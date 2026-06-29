import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

# =========================
# CONFIG
# =========================
input_file = "PRP Sample Jun (2).xlsx"
output_file = "PRP_Final_Output3.xlsx"
sheet_name = "TPRM Web-Portal Export"

# =========================
# LOAD DATA
# =========================
df = pd.read_excel(input_file, sheet_name=sheet_name, engine="openpyxl")

# Clean column names
df.columns = df.columns.str.strip().str.lower()

print("Available columns:", df.columns)

# =========================
# PIVOT TABLE
# =========================
pivot = pd.pivot_table(
    df,
    index="zone_assessing",
    columns="assessment_status",
    values="id",
    aggfunc="count",
    fill_value=0,
    margins=True,
    margins_name="Grand Total"
)

# ✅ Ensure columns exist
for col in ["ACTIVE", "Active", "Deprioritized", "Duplicate"]:
    if col not in pivot.columns:
        pivot[col] = 0

# =========================
# ✅ NEW: COMBINE ACTIVE COLUMNS (IMPORTANT)
# =========================
pivot["Active_Total"] = pivot.get("ACTIVE", 0) + pivot.get("Active", 0)

# =========================
# SUMMARY (for chart 1)
# =========================
summary_df = pd.DataFrame({
    "Status": ["Active", "Deprioritized", "Duplicate"],
    "Count": [
        pivot.loc["Grand Total"]["Active_Total"],
        pivot.loc["Grand Total"].get("Deprioritized", 0),
        pivot.loc["Grand Total"].get("Duplicate", 0)
    ]
})

# =========================
# ✅ NEW: ZONE ACTIVE TABLE (FOR SECOND CHART)
# =========================
zone_active = pivot.loc[pivot.index != "Grand Total"].copy()
zone_active = zone_active.reset_index()

zone_active = zone_active[["zone_assessing", "Active_Total"]]
zone_active.columns = ["Zone", "Active"]

# =========================
# WRITE TO EXCEL
# =========================
with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    pivot.to_excel(writer, sheet_name="Output", startrow=0, startcol=0)
    summary_df.to_excel(writer, sheet_name="Output", startrow=2, startcol=8, index=False)
    zone_active.to_excel(writer, sheet_name="Output", startrow=10, startcol=0, index=False)

# =========================
# LOAD WORKBOOK
# =========================
wb = load_workbook(output_file)
ws = wb["Output"]

# =========================
# CHART 1 (STATUS OVERVIEW)
# =========================
chart = BarChart()
chart.title = "Assessment Status Overview"
chart.y_axis.title = "Count"
chart.x_axis.title = "Assessment Status"

data = Reference(ws, min_col=10, min_row=3, max_row=6)
cats = Reference(ws, min_col=9, min_row=4, max_row=6)

chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

chart.dLbls = DataLabelList()
chart.dLbls.showVal = True

ws.add_chart(chart, "L5")

# =========================
# ✅ NEW: CHART 2 (ACTIVE BY ZONE)
# =========================
chart2 = BarChart()
chart2.title = "Active by Zone"
chart2.y_axis.title = "Count"
chart2.x_axis.title = "Zone"

# dynamic range for zones
start_row = 11
end_row = start_row + len(zone_active)

data2 = Reference(ws, min_col=2, min_row=start_row, max_row=end_row)
cats2 = Reference(ws, min_col=1, min_row=start_row + 1, max_row=end_row)

chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)

chart2.dLbls = DataLabelList()
chart2.dLbls.showVal = True

ws.add_chart(chart2, "L20")

# =========================
# SAVE
# =========================
wb.save(output_file)

print("✅ DONE: Original + Active Zone chart added successfully!")