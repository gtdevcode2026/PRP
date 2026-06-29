import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

# =========================
# CONFIG
# =========================
input_file = "PRP Sample Jun (2).xlsx"
output_file = "PRP_Final_Output2.xlsx"
sheet_name = "TPRM Web-Portal Export"

# =========================
# LOAD DATA
# =========================
df = pd.read_excel(input_file, sheet_name=sheet_name, engine="openpyxl")

# Clean columns
df.columns = df.columns.str.strip().str.lower().str.replace(" ", "_")

# =========================
# MAIN PIVOT (Zone vs Status)
# =========================
pivot = pd.pivot_table(
    df,
    index="zone_assessing",
    columns="assessment_status",
    values="id",
    aggfunc="count",
    fill_value=0,
    margins=True,
    margins_name="Grand Total"
)

# Ensure all columns exist
for col in ["active", "deprioritized", "duplicate"]:
    if col not in pivot.columns:
        pivot[col] = 0

# =========================
# SUMMARY PIVOT (FOR CHART 1)
# =========================
summary_df = pd.DataFrame({
    "Status": ["Active", "Deprioritized", "Duplicate"],
    "Count": [
        pivot.loc["Grand Total"].get("active", 0),
        pivot.loc["Grand Total"].get("deprioritized", 0),
        pivot.loc["Grand Total"].get("duplicate", 0)
    ]
})

# =========================
# ZONE TABLE (Pivot 2 clean format)
# =========================
zone_table = pivot.loc[pivot.index != "Grand Total", ["active", "deprioritized", "duplicate"]]
zone_table = zone_table.reset_index()

zone_table.columns = ["Zone", "Active", "Deprioritized", "Duplicate"]

# =========================
# WRITE TO EXCEL
# =========================
with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    pivot.to_excel(writer, sheet_name="Output", startrow=0, startcol=0)
    summary_df.to_excel(writer, sheet_name="Output", startrow=2, startcol=8, index=False)
    zone_table.to_excel(writer, sheet_name="Output", startrow=10, startcol=0, index=False)

# =========================
# LOAD FILE FOR CHARTS
# =========================
wb = load_workbook(output_file)
ws = wb["Output"]

# =========================
# CHART 1 (STATUS OVERVIEW)
# =========================
chart1 = BarChart()
chart1.title = "Assessment Status Overview"
chart1.y_axis.title = "Count"
chart1.x_axis.title = "Status"

data1 = Reference(ws, min_col=10, min_row=3, max_row=6)
cats1 = Reference(ws, min_col=9, min_row=4, max_row=6)

chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)

chart1.dLbls = DataLabelList()
chart1.dLbls.showVal = True

ws.add_chart(chart1, "L5")

# =========================
# CHART 2 (ZONE ACTIVE BAR - FIXED)
# =========================
chart2 = BarChart()
chart2.title = "Active by Zone"
chart2.y_axis.title = "Count"
chart2.x_axis.title = "Zone"

# ✅ Updated correct range (dynamic based on zone_table)
row_start = 11
row_end = row_start + len(zone_table)

data2 = Reference(ws, min_col=2, min_row=row_start, max_row=row_end)
cats2 = Reference(ws, min_col=1, min_row=row_start + 1, max_row=row_end)

chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)

# ✅ show values
chart2.dLbls = DataLabelList()
chart2.dLbls.showVal = True

ws.add_chart(chart2, "L20")

# =========================
# SAVE
# =========================
wb.save(output_file)

print("✅ DONE: Both pivots + both charts created successfully!")